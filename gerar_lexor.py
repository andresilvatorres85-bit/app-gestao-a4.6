#!/usr/bin/env python3
"""
Converte Controle_LEXOR.xlsx (abas SIOPLEx e Ações) no módulo src/data/lexor.js
usado pela aba LEXOR do aplicativo.

Correções aplicadas na conversão (ver análise de qualidade de dados):
  1. Ação em notação científica (2E+74 -> 2E74), recuperada da Funcional Programática.
  2. Zeros à esquerda em Função (2 díg.), Subfunção (3), Programa (4).
  3. UO obtida da Funcional Programática (não da tabela Ações), que é ambígua
     quando a mesma ação existe em UOs diferentes.
  4. Objeto/Justificativa normalizados (espaços, quebras de linha, CAIXA ALTA).
  5. Tipo de emenda padronizado (27 variações -> 3 categorias + marca de repetida).
"""
import json
import os
import re
import tempfile
import unicodedata
import urllib.request
from pathlib import Path

import openpyxl

# Diretório do próprio script = raiz do repositório.
RAIZ = Path(__file__).resolve().parent

# Fonte da planilha: repositório separado, tratado como fonte única de verdade.
# O workflow diário do GitHub Actions baixa daqui automaticamente. Pode ser
# sobrescrita por variável de ambiente para testes locais.
FONTE_URL = os.environ.get(
    "LEXOR_XLSX_URL",
    "https://raw.githubusercontent.com/andresilvatorres85-bit/"
    "dados-gestao-a4.6/main/Controle_LEXOR.xlsx",
)
# Caminho local opcional: se existir, tem prioridade (conveniência no dev local).
LOCAL_XLSX = os.environ.get("LEXOR_XLSX", str(RAIZ / "Controle_LEXOR.xlsx"))
DESTINO = str(RAIZ / "src" / "data" / "lexor.js")

# Segunda fonte: "Prospecção de Propostas de Emendas.xlsx", no mesmo repositório.
# Dela saem, por Nr Proposta, o Parlamentar, o Partido e os Valores Negociados
# (GND 3, GND 4 e Total). A mesma proposta pode aparecer em várias linhas —
# uma por parlamentar prospectado — e cada linha vira uma proposta no app.
PROSPEC_URL = os.environ.get(
    "PROSPEC_XLSX_URL",
    "https://raw.githubusercontent.com/andresilvatorres85-bit/"
    "dados-gestao-a4.6/main/Prospec%C3%A7%C3%A3o%20de%20Propostas%20de%20Emendas.xlsx",
)
LOCAL_PROSPEC = os.environ.get(
    "PROSPEC_XLSX", str(RAIZ / "Prospecção de Propostas de Emendas.xlsx"))


def obter_planilha():
    """Devolve o caminho do xlsx a processar.

    Prioriza um arquivo local (se existir); caso contrário, baixa da FONTE_URL.
    No CI não há arquivo local, então baixa sempre a versão mais recente do
    repositório de origem.
    """
    if os.path.isfile(LOCAL_XLSX):
        print(f"Planilha local: {LOCAL_XLSX}")
        return LOCAL_XLSX
    print(f"Baixando planilha de: {FONTE_URL}")
    fd, tmp = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    req = urllib.request.Request(FONTE_URL, headers={"User-Agent": "gerar_lexor"})
    with urllib.request.urlopen(req, timeout=60) as r, open(tmp, "wb") as f:
        f.write(r.read())
    return tmp


def obter_prospeccao():
    """Devolve o caminho do xlsx de Prospecção (local, se existir; senão baixa)."""
    if os.path.isfile(LOCAL_PROSPEC):
        print(f"Prospecção local: {LOCAL_PROSPEC}")
        return LOCAL_PROSPEC
    print(f"Baixando Prospecção de: {PROSPEC_URL}")
    fd, tmp = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    req = urllib.request.Request(PROSPEC_URL, headers={"User-Agent": "gerar_lexor"})
    with urllib.request.urlopen(req, timeout=60) as r, open(tmp, "wb") as f:
        f.write(r.read())
    return tmp


def carregar_prospeccao():
    """Lê a Prospecção e devolve { nr: [ {parlamentar, partido, gnd3n, gnd4n,
    totaln}, ... ] }. Uma entrada por linha da planilha — a mesma proposta pode
    ter várias, uma por parlamentar prospectado. Colunas (1-based): Nr Proposta
    (2), Parlamentar (16), Partido (17), Valor Negociado GND 3 (19), GND 4 (20),
    Total (21)."""
    origem = obter_prospeccao()
    wb = openpyxl.load_workbook(origem, data_only=True, read_only=True)
    ws = wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb[wb.sheetnames[0]]
    por_nr = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        nr = limpa(row[1]) if len(row) > 1 else ""
        if not nr:
            continue
        entrada = {
            "parlamentar": limpa(row[15]) if len(row) > 15 else "",
            "partido": (limpa(row[16]).upper() if len(row) > 16 else ""),
            "gnd3n": num(row[18]) if len(row) > 18 else 0,
            "gnd4n": num(row[19]) if len(row) > 19 else 0,
            "totaln": num(row[20]) if len(row) > 20 else 0,
        }
        por_nr.setdefault(nr, []).append(entrada)
    wb.close()
    return por_nr

# Siglas que devem continuar em caixa alta ao converter texto de CAIXA ALTA
SIGLAS = {
    "EB", "ESA", "OM", "TI", "GND", "UO", "CIA", "BC", "BI", "BIB", "RCB", "RCC",
    "SISFRON", "ASTROS", "IMBEL", "EASA", "CMS", "CML", "CMNE", "CMO", "CMA", "CMP",
    "CMSE", "CMN", "COTER", "DECEx", "DEC", "COLOG", "SEF", "DGP", "EME", "GSI",
    "PDF", "GPS", "UTI", "AMAN", "EsPCEx", "IME", "CMB", "PQRMNT", "MEM", "CIGE",
    "SIOP", "CNPJ", "PLOA", "LOA", "LDO", "PPA", "RP", "II", "III", "IV", "V",
    "VI", "VII", "VIII", "IX", "X", "XI", "XII", "1º", "2º", "3º", "4º", "5º",
}
PALAVRAS_MINUSCULAS = {
    "de", "da", "do", "das", "dos", "e", "em", "no", "na", "nos", "nas",
    "a", "o", "as", "os", "para", "com", "por", "ao", "aos", "à", "às", "ou",
}


def limpa(v):
    """Normaliza qualquer célula em string limpa (sem quebras de linha nem espaços duplos)."""
    if v is None:
        return ""
    s = str(v).replace("\r", " ").replace("\n", " ")
    s = re.sub(r"\s+", " ", s).strip()
    return "" if s.lower() in ("none", "nan") else s


def caixa_alta(s):
    letras = [c for c in s if c.isalpha()]
    if not letras:
        return False
    return all(c.isupper() for c in letras)


def frase(s):
    """Converte texto em CAIXA ALTA para caixa de frase, preservando siglas conhecidas."""
    s = limpa(s)
    if not s or not caixa_alta(s):
        return s[:1].upper() + s[1:] if s else s
    saida = []
    for tok in s.split(" "):
        nucleo = tok.strip("().,;:/-")
        if nucleo in SIGLAS or (len(nucleo) <= 4 and nucleo.isupper() and not any(
                c in "AEIOU" for c in nucleo)):
            saida.append(tok)
        else:
            saida.append(tok.lower())
    r = " ".join(saida)
    # primeira letra alfabética em maiúscula
    for i, c in enumerate(r):
        if c.isalpha():
            return r[:i] + c.upper() + r[i + 1:]
    return r


def num(v):
    """Converte célula em inteiro (reais), tolerando texto e vazio."""
    if v is None or v == "":
        return 0
    if isinstance(v, (int, float)):
        return int(round(v))
    s = re.sub(r"[^\d,.-]", "", str(v))
    if not s:
        return 0
    s = s.replace(".", "").replace(",", ".")
    try:
        return int(round(float(s)))
    except ValueError:
        return 0


def cientifica(v):
    """Detecta valores que o Excel corrompeu em notação científica (ex.: 2E74 -> 2e+74)."""
    if isinstance(v, float) and (abs(v) >= 1e15 or "e+" in repr(v).lower()):
        return True
    return bool(re.search(r"\de\+\d", str(v).lower()))


def partes_fp(fp):
    """Quebra a Funcional Programática '10.52921.05.363.6112.21GN.0000'."""
    p = [x.strip() for x in limpa(fp).split(".")]
    if len(p) < 7:
        return {}
    # p[0] é a esfera orçamentária (10 = Fiscal, 20 = Seguridade Social)
    return {"esfera": p[0], "uo": p[1], "funcao": p[2], "subfuncao": p[3],
            "programa": p[4], "acao": p[5].upper(), "subtitulo": p[6]}


def pad(v, n, fp_val=""):
    """Zero-padding, com fallback para o valor vindo da Funcional Programática."""
    s = limpa(v)
    if not s or cientifica(v):
        s = limpa(fp_val)
    if not s:
        return ""
    if s.replace(".", "").isdigit():
        s = s.split(".")[0]
        return s.zfill(n)
    return s


RE_REPETIDA = re.compile(r"repetid[ao]|repetido", re.I)


def classifica_tipo(t):
    """Padroniza as 27 variações da coluna Tipo."""
    t = limpa(t)
    base = unicodedata.normalize("NFKD", t.lower())
    repetida = bool(RE_REPETIDA.search(t))
    if "bancada" in base:
        cat = "Emenda de Bancada"
    elif "comiss" in base:
        cat = "Emenda de Comissão"
    elif "individual" in base:
        cat = "Emenda Individual"
    elif not t:
        cat = ""
    else:
        cat = t
    return cat, repetida


# RP conforme a LDO: 6 = individual impositiva, 7 = bancada, 8 = comissão
RP_POR_TIPO = {"Emenda Individual": "6", "Emenda de Bancada": "7", "Emenda de Comissão": "8"}


def main():
    prospec = carregar_prospeccao()
    origem = obter_planilha()
    wb = openpyxl.load_workbook(origem, data_only=True)

    # ---------- aba Ações: tabela de apoio (o "PROCV" do modelo do Word) ----------
    aba = wb["Ações"]
    acoes = {}
    uos = {}
    for r in range(2, aba.max_row + 1):
        cod = limpa(aba.cell(r, 5).value).upper()
        if not cod:
            continue
        uo_cod = limpa(aba.cell(r, 10).value)
        uo_nome = limpa(aba.cell(r, 11).value)
        if uo_cod:
            uos[uo_cod] = uo_nome
        acoes[cod] = {
            "descricao": limpa(aba.cell(r, 7).value),
            "subtitulo": limpa(aba.cell(r, 6).value) or "XXXX",
            "orgaoCod": limpa(aba.cell(r, 8).value),
            "orgaoNome": limpa(aba.cell(r, 9).value),
            "uoCod": uo_cod,
            "uoNome": uo_nome,
            "produto": limpa(aba.cell(r, 12).value),
            "unidade": limpa(aba.cell(r, 13).value) or "unidade",
            "meta": limpa(aba.cell(r, 14).value) or "1",
            "seq": limpa(aba.cell(r, 15).value),
            "cnpj": limpa(aba.cell(r, 16).value),
        }

    # ---------- aba SIOPLEx: uma proposta por linha ----------
    ws = wb["SIOPLEx"]
    props = []
    for r in range(2, ws.max_row + 1):
        nr = limpa(ws.cell(r, 1).value)
        if not nr:
            continue
        fp_bruta = ws.cell(r, 25).value
        fp = partes_fp(fp_bruta)

        acao_raw = ws.cell(r, 12).value
        acao = limpa(acao_raw).upper()
        if not acao or cientifica(acao_raw):
            acao = fp.get("acao", "")
        if acao.endswith(".0"):
            acao = acao[:-2]

        tipo, repetida = classifica_tipo(ws.cell(r, 2).value)
        uo_cod = fp.get("uo", "") or acoes.get(acao, {}).get("uoCod", "")

        # Valores ORIGINAIS (M/N/O) vêm do Controle_LEXOR. Os NEGOCIADOS e o
        # Parlamentar/Partido passam a vir da Prospecção (ver expansão abaixo).
        gnd3 = num(ws.cell(r, 13).value)
        gnd4 = num(ws.cell(r, 14).value)
        total = num(ws.cell(r, 15).value) or (gnd3 + gnd4)

        registro = {
            "nr": nr,
            "tipo": tipo,
            "rep": 1 if repetida else 0,
            "proponente": limpa(ws.cell(r, 3).value),
            "ods": limpa(ws.cell(r, 4).value),
            "beneficiario": limpa(ws.cell(r, 5).value) or limpa(ws.cell(r, 20).value),
            "cidade": limpa(ws.cell(r, 6).value),
            "uf": limpa(ws.cell(r, 7).value).upper(),
            "objeto": frase(ws.cell(r, 8).value),
            "funcao": pad(ws.cell(r, 9).value, 2, fp.get("funcao")),
            "subfuncao": pad(ws.cell(r, 10).value, 3, fp.get("subfuncao")),
            "programa": pad(ws.cell(r, 11).value, 4, fp.get("programa")),
            "acao": acao,
            "uo": uo_cod,
            "esfera": fp.get("esfera", ""),
            "subtitulo": fp.get("subtitulo") or "XXXX",
            "gnd3": gnd3,
            "gnd4": gnd4,
            "total": total,
            "justificativa": frase(ws.cell(r, 16).value),
            "cmdo": limpa(ws.cell(r, 21).value),
            "catalogo": limpa(ws.cell(r, 22).value),
            "grupo": limpa(ws.cell(r, 23).value),
            "opus": limpa(ws.cell(r, 24).value),
            "fp": limpa(fp_bruta),
            "obs": limpa(ws.cell(r, 26).value),
            "status": limpa(ws.cell(r, 29).value),
            "rp": RP_POR_TIPO.get(tipo, "6"),
            "exportado": limpa(ws.cell(r, 40).value),
        }
        # descarta linhas em branco da planilha (só têm o Nr Proposta preenchido)
        if not registro["beneficiario"] and not registro["objeto"] and not registro["acao"]:
            continue

        # Expansão pela Prospecção: uma proposta por parlamentar prospectado.
        # Sem correspondência, entra uma vez só (sem parlamentar/negociados).
        corresps = prospec.get(nr, [])
        if not corresps:
            corresps = [{}]
        n = len(corresps)
        for k, extra in enumerate(corresps):
            rec = dict(registro)
            rec["parlamentar"] = extra.get("parlamentar", "")
            rec["partido"] = extra.get("partido", "")
            rec["gnd3n"] = extra.get("gnd3n", 0)
            rec["gnd4n"] = extra.get("gnd4n", 0)
            rec["totaln"] = extra.get("totaln", 0)
            # id único: nr quando não duplica; nr#k quando a mesma proposta se
            # repete para vários parlamentares (necessário na seleção/espelho).
            uid = nr if n == 1 else f"{nr}#{k + 1}"
            # omite chaves vazias para reduzir o pacote; uid é sempre mantido.
            enxuto = {kk: vv for kk, vv in rec.items() if vv not in ("", 0)}
            enxuto["uid"] = uid
            props.append(enxuto)

    # ---------- relatório de consistência ----------
    sem_acao = [p["nr"] for p in props if not p.get("acao") or p.get("acao") not in acoes]
    div_uo = [p["nr"] for p in props
              if p.get("acao") in acoes and p.get("uo") and p["uo"] != acoes[p["acao"]]["uoCod"]]
    sem_valor = [p["nr"] for p in props
                 if p.get("gnd3", 0) + p.get("gnd4", 0)
                 + p.get("gnd3n", 0) + p.get("gnd4n", 0) == 0]
    negociados = [p["nr"] for p in props if p.get("gnd3n", 0) or p.get("gnd4n", 0)]
    sem_autor = [p["nr"] for p in props if not p.get("parlamentar")]

    cab = f"""// GERADO AUTOMATICAMENTE a partir de Controle_LEXOR.xlsx — não editar à mão.
// Fonte: repositório dados-gestao-a4.6 (branch main).
// Parlamentar, Partido e os Valores Negociados (GND 3, GND 4 e Total) vêm da
// planilha "Prospecção de Propostas de Emendas.xlsx" (mesmo repositório), por
// Nr Proposta. Quando a mesma proposta foi prospectada para vários
// parlamentares, ela é duplicada — uma linha por parlamentar (campo `uid`).
// Atualização automática: workflow diário do GitHub Actions (06:00 BRT) baixa a
// planilha, roda gerar_lexor.py e publica. Para atualizar na hora, dispare o
// workflow "Deploy no GitHub Pages" manualmente (workflow_dispatch).
//
// Propostas: {len(props)}   |   Ações cadastradas: {len(acoes)}
// Correções aplicadas na conversão:
//   · Ação em notação científica recuperada da Funcional Programática ({len([p for p in props if p.get("acao") == "2E74"])} casos de 2E74)
//   · Zeros à esquerda restaurados em Função / Subfunção / Programa
//   · UO lida da Funcional Programática (a tabela Ações é ambígua por ação)
//   · Objeto e Justificativa convertidos de CAIXA ALTA para caixa de frase
//   · Tipo de emenda padronizado em 3 categorias (+ marca de repetida)
//
// Pendências de preenchimento detectadas na origem:
//   · {len(sem_acao)} propostas sem ação correspondente na aba Ações
//   · {len(div_uo)} propostas com UO divergente entre a Funcional Programática e a aba Ações
//   · {len(sem_valor)} propostas sem valor em GND 3 e GND 4
//   · {len(negociados)} propostas com valor negociado (prevalece sobre o valor original)
//   · {len(sem_autor)} propostas sem parlamentar autor definido
"""

    js = cab + "\nexport const ACOES_LEXOR = " + json.dumps(acoes, ensure_ascii=False, indent=0) + ";\n"
    js += "\nexport const UO_LEXOR = " + json.dumps(uos, ensure_ascii=False, indent=0) + ";\n"
    js += "\nexport const PROPOSTAS_LEXOR = " + json.dumps(props, ensure_ascii=False, separators=(",", ":")) + ";\n"

    with open(DESTINO, "w", encoding="utf-8") as f:
        f.write(js)

    print(f"OK -> {DESTINO}")
    print(f"  propostas: {len(props)} | ações: {len(acoes)} | UOs: {len(uos)}")
    print(f"  sem ação na tabela: {len(sem_acao)}")
    print(f"  UO divergente:      {len(div_uo)}")
    print(f"  sem valor:          {len(sem_valor)}")
    print(f"  sem autor:          {len(sem_autor)}")


if __name__ == "__main__":
    main()
